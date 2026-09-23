#!/usr/bin/env python3
"""Build DOCX editions of the 16 standalone theorem monographs (RU/EN).

Pipeline per document:
  1. preprocess LaTeX  — tcolorbox envs -> blockquote with bold lead-in,
                          strip \\hrule rules / \\vspace / \\vfill,
                          babel guillemets <<..>> -> «..»
  2. pandoc            — LaTeX -> DOCX with native OMML equations,
                         styled via a customized reference.docx
  3. python-docx post  — split/center the title block, style the footer,
                         set core properties (title/author/subject/lang)
  4. validate          — oMath count, leftovers, PDF integrity (pypdf)
"""
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = Path(__file__).resolve().parent.parent
THEOREMS = REPO / "theorems"
TMP = Path("/tmp/docx_build")
REF = TMP / "ref_styled.docx"
SERIF = "Times New Roman"
ACCENT = "1F4E79"

BOXES = {
    "ru": {"statusbox": "Сводка доказанного",
           "databox": "Данные протокола",
           "verifybox": "Верификация"},
    "en": {"statusbox": "Summary of results",
           "databox": "Protocol data",
           "verifybox": "Verification"},
}
LANGMETA = {
    "ru": {"lang": "ru-RU", "author": "Исаев Исхак Хамзатович", "label": "RU"},
    "en": {"lang": "en-US", "author": "Isaev Iskhak Khamzatovich", "label": "EN"},
}


# ---------------------------------------------------------------- preprocess
def preprocess(tex: str, lang: str) -> str:
    t = tex
    for env, title in BOXES[lang].items():
        def repl(m, title=title):
            opt = (m.group(1) or "").strip()
            head = f"{title} ({opt})" if opt else title
            return "\\begin{quote}\n\\textbf{" + head + ".} "
        t = re.sub(r"\\begin\{" + env + r"\}(?:\[([^\]]*)\])?", repl, t)
        t = t.replace("\\end{" + env + "}", "\\end{quote}")
    # decorative color rules -> drop (leaked "height 0.8pt" otherwise)
    t = re.sub(r"\{\\color\{linkblue\}\\hrule\s*height\s*[0-9.]+pt\}", "", t)
    # vertical spacing / fill -> meaningless in DOCX
    t = re.sub(r"\\vspace\*?\{[^}]*\}", "", t)
    t = t.replace("\\vfill", "")
    # babel russian shorthand -> real guillemets (as rendered in the PDFs)
    t = re.sub(r"<<([^<>\n]{1,120})>>", "«\\1»", t)
    return t


# ------------------------------------------------------------- reference docx
def _set_style_font(style, name=SERIF, size=None, bold=None, color=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor.from_string(color)


def _insert_ppr_ordered(ppr, el):
    """Insert el into pPr respecting CT_PPr child order (before spacing/ind/jc)."""
    after_tags = {qn("w:spacing"), qn("w:ind"), qn("w:jc"),
                  qn("w:contextualSpacing"), qn("w:rPr")}
    for child in ppr:
        if child.tag in after_tags:
            child.addprevious(el)
            return
    ppr.append(el)


def make_reference() -> Path:
    TMP.mkdir(parents=True, exist_ok=True)
    raw = TMP / "ref_default.docx"
    if not raw.exists():
        data = subprocess.run(
            ["pandoc", "--print-default-data-file", "reference.docx"],
            check=True, capture_output=True).stdout
        raw.write_bytes(data)
    d = docx.Document(str(raw))
    st = d.styles

    normal = st["Normal"]
    _set_style_font(normal, SERIF, 11)
    pf = normal.paragraph_format
    pf.line_spacing = 1.3
    pf.space_after = Pt(6)

    for name in ("Body Text", "First Paragraph"):
        s = st[name]
        _set_style_font(s, SERIF, 11)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        s.paragraph_format.line_spacing = 1.3

    head_specs = {"Heading 1": 16, "Heading 2": 13, "Heading 3": 12,
                  "Heading 4": 11}
    for name, size in head_specs.items():
        try:
            s = st[name]
        except KeyError:
            continue
        _set_style_font(s, SERIF, size, bold=True, color=ACCENT)
        s.paragraph_format.space_before = Pt(14 if size >= 16 else 10)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True

    # Block Text == pandoc blockquote == our three protocol boxes
    bt = st["Block Text"]
    _set_style_font(bt, SERIF, 10.5)
    bpf = bt.paragraph_format
    bpf.left_indent = Cm(0.45)
    bpf.right_indent = Cm(0.15)
    bpf.space_before = Pt(8)
    bpf.space_after = Pt(8)
    bpf.line_spacing = 1.25
    bpf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ppr = bt.element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), ACCENT)
    bdr.append(left)
    _insert_ppr_ordered(ppr, bdr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F4F6F9")
    _insert_ppr_ordered(ppr, shd)

    for name in ("Caption", "Table Caption", "Image Caption"):
        try:
            s = st[name]
        except KeyError:
            continue
        _set_style_font(s, SERIF, 9.5, color="555555")
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for sec in d.sections:
        sec.top_margin = Cm(2.4)
        sec.bottom_margin = Cm(2.4)
        sec.left_margin = Cm(2.4)
        sec.right_margin = Cm(2.4)
    d.save(str(REF))
    return REF


# ---------------------------------------------------------------- postprocess
def para_lines(p) -> list:
    lines, cur = [], []
    for node in p._p.iter():
        if node.tag == qn("w:t"):
            cur.append(node.text or "")
        elif node.tag == qn("w:br"):
            lines.append("".join(cur))
            cur = []
        elif node.tag == qn("w:tab"):
            cur.append(" ")
    lines.append("".join(cur))
    return [ln.strip() for ln in lines if ln.strip()]


def add_centered_line(anchor, text, size, bold=False, color=None, italic=False):
    p = anchor.insert_paragraph_before("")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def post_process(path: Path, meta: dict):
    d = docx.Document(str(path))
    cp = d.core_properties
    cp.author = meta["author"]
    cp.title = meta["title"]
    cp.subject = meta["subject"]
    cp.language = meta["lang"]
    cp.keywords = "hodge-laboratory, theorem monograph, " + meta["label"]

    paras = d.paragraphs
    h1 = next((i for i, p in enumerate(paras)
               if (p.style.name or "").startswith("Heading")), None)
    if h1 and h1 > 0:
        title_par = paras[0]
        lines = para_lines(title_par)
        if lines:
            sizes = [(20, True, ACCENT, False), (13, False, "404040", False),
                     (11.5, True, "202020", False)]
            anchor = title_par
            for i, line in enumerate(lines):
                if i < len(sizes):
                    sz, bd, col, it = sizes[i]
                else:
                    sz, bd, col, it = 9.5, False, "666666", False
                add_centered_line(anchor, line, sz, bd, col, it)
        # remove the merged original paragraph
        title_par._p.getparent().remove(title_par._p)

    # footer license block -> small gray italic
    nonempty = [p for p in d.paragraphs if p.text.strip()]
    if nonempty:
        tail = nonempty[-1]
        if len(tail.text) > 80:  # the license/verification block
            for r in tail.runs:
                r.font.size = Pt(9)
                r.font.italic = True
                r.font.color.rgb = RGBColor.from_string("707070")
    d.save(str(path))


# ------------------------------------------------------------------- convert
def convert_one(texpath: Path, outpath: Path, lang: str) -> tuple:
    meta = LANGMETA[lang]
    pre = preprocess(texpath.read_text(encoding="utf-8"), lang)
    m = re.search(r"\\LARGE\s*\\bfseries\s*([^{}]+?)\}", pre)
    title = m.group(1).strip() if m else texpath.parent.parent.parent.name
    k = re.search(r"(\d+)\s*/\s*16", pre)
    subject = (f"Standalone theorem monograph {k.group(1)}/16"
               if k else "Standalone theorem monograph")
    meta = dict(meta, title=title, subject=subject)

    tmp_tex = TMP / f"{outpath.stem}.tex"
    TMP.mkdir(parents=True, exist_ok=True)
    tmp_tex.write_text(pre, encoding="utf-8")

    cmd = ["pandoc", "-f", "latex", "-t", "docx",
           f"--reference-doc={REF}", "-M", f"lang={meta['lang']}",
           "-o", str(outpath), str(tmp_tex)]
    r = subprocess.run(cmd, capture_output=True, text=True)
    return r, dict(meta, out=outpath)


# ----------------------------------------------------------------- validate
def validate() -> int:
    problems = 0
    pdfs = sorted(THEOREMS.glob("T*/*.pdf"))
    docxs = sorted(THEOREMS.glob("T*/monograph_*.docx"))
    print(f"\n=== VALIDATION: {len(pdfs)} PDFs, {len(docxs)} DOCX ===")
    try:
        from pypdf import PdfReader
        for pdf in pdfs:
            try:
                n = len(PdfReader(str(pdf)).pages)
                if n < 1:
                    print(f"  ⚠ {pdf.name} has only {n} pages"); problems += 1
            except Exception as e:  # noqa: BLE001
                print(f"  ✗ BROKEN PDF {pdf}: {e}"); problems += 1
    except ImportError:
        print("  (pypdf missing — PDF integrity not re-checked)")
    for dx in docxs:
        try:
            with zipfile.ZipFile(dx) as z:
                xml = z.read("word/document.xml").decode("utf-8")
            n_math = xml.count("<m:oMath")
            text = re.sub(r"<[^>]+>", "", xml)
            bad = [s for s in ("statusbox", "databox", "verifybox",
                               "tcolorbox", "\\begin", "height 0.8")
                   if s in text]
            if len(text) < 1500:
                print(f"  ⚠ {dx}: text too short ({len(text)} chars)")
                problems += 1
            if bad:
                print(f"  ⚠ {dx}: leftovers {bad}")
                problems += 1
            kb = dx.stat().st_size // 1024
            print(f"  ✓ {dx.parent.name}/{dx.name}: {kb} KB, {n_math} equations")
        except Exception as e:  # noqa: BLE001
            print(f"  ✗ BROKEN DOCX {dx}: {e}")
            problems += 1
    return problems


def main():
    only = sys.argv[1] if len(sys.argv) > 1 else None
    make_reference()
    built = 0
    for tdir in sorted(THEOREMS.glob("T*")):
        if not tdir.is_dir():
            continue
        for lang in ("ru", "en"):
            tex = tdir / "latex" / lang / "monograph.tex"
            if not tex.exists():
                print(f"  !! missing {tex}")
                continue
            out = tdir / f"monograph_{LANGMETA[lang]['label']}.docx"
            if only and only not in str(out):
                continue
            r, meta = convert_one(tex, out, lang)
            if r.returncode != 0:
                print(f"  ✗ pandoc FAIL {out}: {r.stderr[:400]}")
                continue
            if r.stderr.strip():
                print(f"  · pandoc note {out.stem}: {r.stderr.strip()[:200]}")
            post_process(out, meta)
            kb = out.stat().st_size // 1024
            print(f"  ✓ built {out.parent.name}/{out.name}: {kb} KB — {meta['title']}")
            built += 1
    print(f"\nBuilt {built} DOCX files.")
    if not only:
        sys.exit(1 if validate() else 0)


if __name__ == "__main__":
    main()
